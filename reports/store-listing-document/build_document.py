import json
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
OUT = ROOT / 'docs' / 'play-store'
ORDER = ['en-US', 'fr-FR', 'de-DE', 'pt-BR', 'zh-CN', 'es-ES', 'hi-IN', 'ja-JP', 'sw', 'am', 'id', 'fil', 'ko-KR', 'ru-RU', 'ar', 'fa']
FONTS = {'zh-CN': 'Microsoft YaHei', 'ja-JP': 'Yu Gothic', 'ko-KR': 'Malgun Gothic', 'hi-IN': 'Nirmala UI', 'am': 'Ebrima'}

entries = [json.loads((WORK / 'source-en.json').read_text(encoding='utf-8-sig'))]
for filename in ['translations-root.json', 'translations-west.json', 'translations-asian.json', 'translations-rtl-sw.json']:
    entries.extend(json.loads((WORK / filename).read_text(encoding='utf-8-sig')))
by_locale = {item['locale']: item for item in entries}
assert len(entries) == len(by_locale) == 16
assert set(by_locale) == set(ORDER)
entries = [by_locale[locale] for locale in ORDER]
for item in entries:
    assert len(item['paragraphs']) == 8
    item['short'] = unicodedata.normalize('NFC', item['short'])
    item['paragraphs'] = [unicodedata.normalize('NFC', p) for p in item['paragraphs']]
    item['full'] = '\n\n'.join(item['paragraphs'])
    item['shortCount'] = len(item['short'])
    item['fullCount'] = len(item['full'])
    assert 1 <= item['shortCount'] <= 80, item['locale']
    assert 1 <= item['fullCount'] <= 4000, item['locale']
    assert len(item['short'].encode('utf-16-le')) // 2 <= 80
    assert len(item['full'].encode('utf-16-le')) // 2 <= 4000
    assert '\n' not in item['short'] and item['short'] == item['short'].strip()
    assert '\ufffd' not in item['full']
    assert 'Qetoret' in item['full'] and '7.0' in item['full'] and '16' in item['full']
    assert set(re.findall(r'https://[^\s]+', item['full'])) == {
        'https://qetoret.com', 'https://qetoret.com/privacy.html', 'https://qetoret.com/delete-account.html'
    }

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.75)
section.left_margin = section.right_margin = Cm(1.9)
section.header_distance = section.footer_distance = Cm(0.7)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.05
    rpr = style.element.get_or_add_rPr()
    color = rpr.find(qn('w:color'))
    if color is not None:
        for attr in ['themeColor', 'themeTint', 'themeShade']:
            color.attrib.pop(qn('w:' + attr), None)
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme']:
            fonts.attrib.pop(qn('w:' + attr), None)
    borders = style.element.find('.//' + qn('w:pBdr'))
    if borders is not None:
        borders.getparent().remove(borders)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)

def font_run(run, locale='en-US', rtl=False):
    font = FONTS.get(locale, 'Arial')
    run.font.name = font
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + attr), font)
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), locale)
    lang.set(qn('w:eastAsia'), locale)
    lang.set(qn('w:bidi'), locale)
    rpr.append(lang)
    if rtl:
        direction = OxmlElement('w:rtl')
        direction.set(qn('w:val'), '1')
        rpr.append(direction)

def direct_text_format(paragraph):
    size = {'Title': 24, 'Subtitle': 12, 'Heading 1': 18, 'Heading 2': 11}.get(paragraph.style.name, 11)
    for run in paragraph.runs:
        if not run.font.name:
            font_run(run)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt(size)
        rpr = run._element.get_or_add_rPr()
        color = rpr.find(qn('w:color'))
        if color is not None:
            for attr in ['themeColor', 'themeTint', 'themeShade']:
                color.attrib.pop(qn('w:' + attr), None)
        cs_size = OxmlElement('w:szCs')
        cs_size.set(qn('w:val'), str(size * 2))
        rpr.append(cs_size)
        if paragraph.style.name in ['Heading 1', 'Heading 2']:
            run.bold = True

def localized_paragraph(text, locale):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if locale in ['ar', 'fa']:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        bidi = OxmlElement('w:bidi')
        p._p.get_or_add_pPr().append(bidi)
        for part in re.split(r'([A-Za-z0-9][A-Za-z0-9 .:/_-]*)', text):
            if part:
                run = p.add_run(part)
                font_run(run, locale, rtl=not bool(re.fullmatch(r'[A-Za-z0-9 .:/_-]+', part)))
    else:
        font_run(p.add_run(text), locale)
    return p

doc.add_paragraph('Qetoret Google Play Store Descriptions', 'Title')
doc.add_paragraph('Short and full descriptions for 16 supported languages', 'Subtitle')
doc.add_paragraph('Each language has a short description for the 80-character field and a full description for the 4000-character field. All entries meet both limits. App name for every language: Qetoret.')
doc.add_paragraph('Copy only the description text into its matching field. The language labels and character counts are for reference. Counts include spaces, punctuation, URLs and paragraph breaks. English United States is the default locale.')
doc.add_paragraph('Character counts', 'Heading 2')
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(6.9), Cm(3.0), Cm(3.1), Cm(3.2)]
for cell, label, width in zip(table.rows[0].cells, ['Language', 'Locale', 'Short of 80', 'Full of 4000'], widths):
    cell.width = width
    cell.text = label
header = OxmlElement('w:tblHeader')
table.rows[0]._tr.get_or_add_trPr().append(header)
for item in entries:
    row = table.add_row()
    for cell, value, width in zip(row.cells, [item['language'], item['locale'], str(item['shortCount']), str(item['fullCount'])], widths):
        cell.width = width
        cell.text = value
for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ['top', 'left', 'bottom', 'right']:
            el = OxmlElement('w:' + edge)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                el.set(qn('w:' + key), value)
            borders.append(el)
        tcpr.append(borders)
        margins = OxmlElement('w:tcMar')
        for edge in ['top', 'bottom', 'left', 'right']:
            el = OxmlElement('w:' + edge)
            el.set(qn('w:w'), '80')
            el.set(qn('w:type'), 'dxa')
            margins.append(el)
        tcpr.append(margins)
        if i == 0:
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8E8E8')
            tcpr.append(shade)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            for run in p.runs:
                font_run(run)
                run.font.size = Pt(9)
                run.bold = i == 0
doc.add_paragraph('Prepared 8 October 2026. Store locale fil corresponds to the Tagalog interface language. Portuguese uses the Brazil locale and Chinese uses Simplified Chinese.').paragraph_format.space_before = Pt(8)

for item in entries:
    doc.add_page_break()
    heading = doc.add_paragraph(re.sub(r'[()]', '', item['language']), 'Heading 1')
    p = doc.add_paragraph('Google Play locale ' + item['locale'] + '   App name Qetoret')
    p.paragraph_format.space_after = Pt(3)
    localized_paragraph(item['nativeLanguage'], item['locale']).paragraph_format.space_after = Pt(8)
    doc.add_paragraph('Short description', 'Heading 2')
    doc.add_paragraph(f"{item['shortCount']} of 80 characters").paragraph_format.space_after = Pt(5)
    localized_paragraph(item['short'], item['locale'])
    doc.add_paragraph('Full description', 'Heading 2')
    doc.add_paragraph(f"{item['fullCount']} of 4000 characters").paragraph_format.space_after = Pt(5)
    for text in item['paragraphs']:
        localized_paragraph(text, item['locale'])

for paragraph in doc.paragraphs:
    direct_text_format(paragraph)

doc.core_properties.title = 'Qetoret Google Play Store Descriptions'
doc.core_properties.subject = 'Localized short and full store descriptions in 16 languages'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.keywords = 'Qetoret, Google Play, store descriptions, localization'
docx_path = OUT / 'qetoret-google-play-descriptions-16-languages.docx'
doc.save(docx_path)

# Confirm that the saved DOCX contains every exact field string in logical order.
reloaded = Document(docx_path)
saved_paragraphs = [p.text for p in reloaded.paragraphs]
for item in entries:
    assert item['short'] in saved_paragraphs, item['locale']
    position = saved_paragraphs.index(item['short'])
    full_start = position + 3
    assert saved_paragraphs[full_start:full_start + 8] == item['paragraphs'], item['locale']

text_parts = ['Qetoret Google Play Store Descriptions', 'App name for every language: Qetoret', 'Copy only the text under each description label. Character counts include spaces and paragraph breaks.']
for item in entries:
    text_parts.extend([
        f"{item['language']} | {item['locale']}",
        f"SHORT DESCRIPTION ({item['shortCount']}/80)", item['short'],
        f"FULL DESCRIPTION ({item['fullCount']}/4000)", item['full'],
    ])
txt_path = OUT / 'qetoret-google-play-descriptions-16-languages.txt'
txt_path.write_text('\n\n'.join(text_parts) + '\n', encoding='utf-8')
counts = [{k: item[k] for k in ['locale', 'language', 'shortCount', 'fullCount']} for item in entries]
(WORK / 'character-counts.json').write_text(json.dumps({'unicodeNormalization': 'NFC', 'countIncludes': 'spaces punctuation URLs and LF paragraph separators', 'locales': counts}, ensure_ascii=False, indent=2), encoding='utf-8')
(WORK / 'compiled-listings.json').write_text(json.dumps(entries, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'docx': str(docx_path), 'text': str(txt_path), 'languages': len(entries), 'counts': counts}, ensure_ascii=False))
